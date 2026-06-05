from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "AI面试服务接口对齐文档-供对方调用版.docx"


def font(run, size=10.5, bold=False, color=None, name="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def p(doc, text=""):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(6)
    run = para.add_run(text)
    font(run)
    return para


def h(doc, text, level=1):
    para = doc.add_heading(text, level=level)
    for run in para.runs:
        font(run, size=14 if level == 1 else 12, bold=True, color=(31, 78, 121) if level == 1 else (47, 84, 150))
    return para


def code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, "F7F7F7")
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(3)
    lines = text.strip("\n").splitlines()
    for i, line in enumerate(lines):
        if i:
            para.add_run().add_break()
        run = para.add_run(line)
        font(run, size=8.5, name="Consolas")
    doc.add_paragraph()


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = t.rows[0].cells[idx]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        font(run, size=9, bold=True)
        shade(cell, "EAF2F8")
    for row in rows:
        cells = t.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = ""
            run = cells[idx].paragraphs[0].add_run(str(value))
            font(run, size=9)
    doc.add_paragraph()
    return t


def bullets(doc, items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        run = para.add_run(item)
        font(run)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)
    doc.styles["Normal"].font.name = "Microsoft YaHei"
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("AI 面试任务创建与结果回调接口对齐文档")
    font(r, size=20, bold=True, color=(31, 78, 121))
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("严格按合作方 Word 接口协议对齐版")
    font(r, size=11, color=(90, 90, 90))

    h(doc, "1. 对接结论")
    p(doc, "本对接方案严格以合作方提供的《AI 面试任务创建与结果回调 API 设计文档》为主协议。合作方请求字段名、层级和响应字段名保持不变；我方不要求合作方改成我方内部字段。")
    p(doc, "我方在接口入口增加一层字段映射/适配层：例如合作方传 job.title，我方内部映射为 job_name；合作方传 job.description，我方内部映射为 job_description；合作方传 resume.file_url，我方下载简历并解析为内部 resume_json。")
    p(doc, "我方内部有但合作方接口不需要的字段，不要求合作方传入；必要时由我方使用默认值、派生值或内部配置补齐。只有确实影响传输和业务处理的类型差异，才在本文档中列为建议调整项。")

    h(doc, "2. 必须保持不变的合作方字段")
    table(doc, ["字段", "是否保持", "说明"], [
        ("request_id", "保持不变", "作为创建任务幂等键"),
        ("scene", "保持不变", "固定值 ai_interview"),
        ("candidate.id/name/phone/email", "保持不变", "合作方候选人对象结构不变"),
        ("job.id/title/description", "保持不变", "合作方岗位对象结构不变"),
        ("resume.file_name/file_url", "保持不变", "合作方只提供文件名和下载地址"),
        ("callback_url", "保持不变", "合作方提供给我方回调结果的地址"),
        ("external_session_id", "保持不变", "我方响应中按合作方协议返回"),
        ("session_url", "保持不变", "我方响应中按合作方协议返回"),
        ("expires_at", "保持不变", "我方响应中按合作方协议返回"),
    ])

    h(doc, "3. 接口总览")
    table(doc, ["接口", "方法", "路径", "说明"], [
        ("创建 AI 面试任务", "POST", "/api/external/ai-session/create", "合作方调用我方，创建面试任务并获取面试链接"),
        ("面试结果回调", "POST", "合作方在创建任务时传入的 callback_url", "候选人完成面试后，我方主动向该地址 POST 面试结果"),
        ("面试链接访问", "GET", "创建接口返回的 session_url", "候选人打开链接进入我方面试流程"),
    ])

    h(doc, "4. 通用约定")
    table(doc, ["项目", "约定"], [
        ("协议", "HTTPS，测试环境可使用 HTTP"),
        ("请求格式", "application/json"),
        ("响应格式", "application/json"),
        ("字符编码", "UTF-8"),
        ("鉴权", "Authorization: Bearer {token}"),
        ("幂等", "创建任务按 request_id 幂等，结果回调按 request_id + external_session_id 幂等"),
        ("时间格式", "ISO 8601，例如 2026-06-02T15:30:00+08:00"),
    ])

    h(doc, "5. 创建 AI 面试任务")
    table(doc, ["项目", "内容"], [
        ("请求方式", "POST"),
        ("请求路径", "/api/external/ai-session/create"),
        ("Content-Type", "application/json"),
        ("Authorization", "Bearer token，由双方线下配置"),
    ])
    code(doc, """
{
  "request_id": "7e8f4d4a-246a-4d94-ae8b-a7b7f4d0a001",
  "scene": "ai_interview",
  "candidate": {
    "id": 123,
    "name": "周文胜",
    "phone": "18807372758",
    "email": "a18807372758@gmail.com"
  },
  "job": {
    "id": 456,
    "title": "产品经理",
    "description": "岗位详情，包含职位描述、岗位职责、任职要求、加分项等完整信息。"
  },
  "resume": {
    "file_name": "周文胜-产品经理.pdf",
    "file_url": "https://partner-domain.com/files/resumes/xxx.pdf"
  },
  "callback_url": "https://partner-domain.com/api/integrations/ai-session/callback/"
}
""")
    table(doc, ["字段", "类型", "必填", "说明"], [
        ("request_id", "string", "是", "合作方生成的唯一请求 ID，用于防重复创建"),
        ("scene", "string", "是", "固定值：ai_interview"),
        ("candidate", "object", "是", "候选人信息"),
        ("candidate.id", "number/string", "是", "候选人在合作方系统中的 ID"),
        ("candidate.name", "string", "是", "候选人姓名"),
        ("candidate.phone", "string", "否", "候选人手机号"),
        ("candidate.email", "string", "是", "候选人邮箱，用于身份校验"),
        ("job", "object", "是", "岗位信息"),
        ("job.id", "number/string", "是", "岗位在合作方系统中的 ID"),
        ("job.title", "string", "是", "岗位名称"),
        ("job.description", "string", "是", "岗位详情/JD，用于生成 AI 面试问题"),
        ("resume", "object", "是", "简历文件信息"),
        ("resume.file_name", "string", "是", "简历文件名"),
        ("resume.file_url", "string", "是", "合作方提供给我方的简历下载地址；我方通过该地址下载并解析简历"),
        ("callback_url", "string", "是", "合作方提供给我方的结果接收地址；面试完成后由我方主动 POST 回调该地址"),
    ])

    h(doc, "6. 创建任务成功响应")
    code(doc, """
{
  "success": true,
  "external_session_id": "ai_4089dfa7fbdb42ce",
  "session_url": "https://our-domain.com/api/interview/welcome/?config_id=4089dfa7fbdb42ce",
  "expires_at": "2026-06-09T10:00:00+08:00"
}
""")
    table(doc, ["字段", "类型", "说明"], [
        ("success", "boolean", "是否创建成功"),
        ("external_session_id", "string", "我方 AI 面试任务 ID；合作方后续用它识别任务"),
        ("session_url", "string", "候选人访问 AI 面试的链接，合作方发送给候选人"),
        ("expires_at", "string", "面试链接过期时间"),
    ])

    h(doc, "7. 创建任务失败响应")
    code(doc, """
{
  "success": false,
  "error_code": "INVALID_PARAM",
  "error_message": "job.description不能为空"
}
""")
    table(doc, ["错误码", "说明"], [
        ("INVALID_PARAM", "请求参数错误"),
        ("UNAUTHORIZED", "鉴权失败"),
        ("FORBIDDEN", "无权限访问"),
        ("DUPLICATE_REQUEST", "request_id 已存在但无法返回已创建任务"),
        ("RESUME_DOWNLOAD_FAILED", "简历文件无法访问或下载失败"),
        ("JOB_INFO_INVALID", "岗位信息不完整"),
        ("CANDIDATE_INFO_INVALID", "候选人信息不完整"),
        ("RATE_LIMITED", "请求频率超限"),
        ("INTERNAL_ERROR", "系统内部错误"),
        ("SERVICE_UNAVAILABLE", "服务暂时不可用"),
    ])

    h(doc, "8. 创建任务幂等规则")
    bullets(doc, [
        "合作方每次创建任务必须携带唯一 request_id。",
        "我方首次收到 request_id 时创建 AI 面试任务。",
        "如果相同 request_id 再次请求，且请求内容一致，我方直接返回已创建的 external_session_id、session_url、expires_at。",
        "如果相同 request_id 对应的核心内容不一致，我方返回 INVALID_PARAM 或 DUPLICATE_REQUEST，避免生成多个任务。",
    ])

    h(doc, "9. 面试结果回调")
    p(doc, "候选人完成面试后，我方根据创建任务时合作方传入的 callback_url 主动回调合作方。注意：callback_url 不是合作方调用我方的接口，而是我方调用合作方的结果接收接口。合作方应在 0.5 秒内返回确认结果。")
    code(doc, """
{
  "request_id": "7e8f4d4a-246a-4d94-ae8b-a7b7f4d0a001",
  "external_session_id": "ai_4089dfa7fbdb42ce",
  "status": "completed",
  "completed_at": "2026-06-02T15:30:00+08:00",
  "result": {
    "score": 86,
    "summary_for_hr": "候选人表达清晰，项目经历与岗位匹配度较高，建议进入下一轮。",
    "summary_for_candidate": "你在项目表达和技术思路方面表现较好，可继续补充业务结果和量化指标。",
    "dimensions": [
      {
        "name": "项目经验",
        "score": 88,
        "comment": "能说明项目职责、技术方案和优化结果"
      },
      {
        "name": "沟通表达",
        "score": 84,
        "comment": "表达较清楚，结构性较好"
      }
    ],
    "transcript_url": "https://our-domain.com/api/ai-interview/reports/ai_4089dfa7fbdb42ce/transcript/",
    "report_url": "https://our-domain.com/api/ai-interview/reports/ai_4089dfa7fbdb42ce/"
  }
}
""")
    table(doc, ["字段", "类型", "必填", "说明"], [
        ("request_id", "string", "是", "创建任务时合作方传入的请求 ID"),
        ("external_session_id", "string", "是", "我方 AI 面试任务 ID"),
        ("status", "string", "是", "面试状态：completed、failed、expired、cancelled 等"),
        ("completed_at", "string", "是", "面试完成时间"),
        ("result", "object", "是", "面试结果；failed/expired/cancelled 时可为空或只返回说明"),
        ("result.score", "number", "是", "综合得分，0-100"),
        ("result.summary_for_hr", "string", "是", "给 HR 看的总结"),
        ("result.summary_for_candidate", "string", "否", "给候选人看的反馈摘要"),
        ("result.dimensions", "array", "是", "分维度评分"),
        ("result.dimensions[].name", "string", "是", "维度名称"),
        ("result.dimensions[].score", "number", "是", "维度得分，0-100"),
        ("result.dimensions[].comment", "string", "是", "维度评价"),
        ("result.transcript_url", "string", "否", "面试问答记录地址"),
        ("result.report_url", "string", "否", "面试报告地址"),
    ])

    h(doc, "10. 合作方回调确认响应")
    p(doc, "合作方收到我方回调后，只需要尽快返回确认响应。只要业务侧已接收或已处理过相同结果，都应返回 received=true，避免我方继续重试。")
    code(doc, """
{
  "received": true,
  "request_id": "7e8f4d4a-246a-4d94-ae8b-a7b7f4d0a001"
}
""")
    p(doc, "如果合作方返回 received 不为 true，或 0.5 秒内未响应，我方可按重试策略再次发送相同回调。重复回调时 request_id、external_session_id、status、result 内容保持一致。")
    code(doc, """
{
  "received": false,
  "error_code": "INVALID_SESSION",
  "error_message": "external_session_id不存在"
}
""")
    table(doc, ["错误码", "说明"], [
        ("INVALID_SESSION", "面试任务不存在"),
        ("INVALID_REQUEST", "请求参数错误"),
        ("DUPLICATE_CALLBACK", "重复回调；合作方也可直接返回 received=true"),
        ("INTERNAL_ERROR", "系统内部错误"),
    ])

    h(doc, "11. 回调重试机制")
    bullets(doc, [
        "合作方应在 0.5 秒内返回 received=true。",
        "如未收到确认，或响应中 received 不为 true，我方每 5 秒重试一次相同结果回调，直到收到 received=true。",
        "重试内容保持完全一致，合作方应基于 request_id 或 external_session_id 幂等处理。",
        "同一任务结果已处理成功后，再次收到相同回调时，合作方仍应返回 received=true。",
    ])

    h(doc, "12. 回调方向说明")
    table(doc, ["字段/动作", "提供方", "调用方", "说明"], [
        ("resume.file_url", "合作方", "我方下载", "合作方在创建任务请求中提供简历下载链接，我方用它下载并解析简历"),
        ("callback_url", "合作方", "我方回调", "合作方在创建任务请求中提供结果接收地址，我方面试完成后 POST 结果到该地址"),
        ("session_url", "我方", "候选人访问", "我方创建任务成功后返回面试访问链接，合作方发送给候选人"),
        ("external_session_id", "我方", "双方保存", "双方后续用该 ID 对账、查询、处理回调幂等"),
    ])

    h(doc, "13. 状态枚举")
    table(doc, ["status", "说明"], [
        ("pending", "待开始"),
        ("in_progress", "面试进行中"),
        ("completed", "已完成"),
        ("failed", "面试失败"),
        ("expired", "链接已过期"),
        ("cancelled", "已取消"),
    ])

    h(doc, "14. 简历文件约定")
    bullets(doc, [
        "resume.file_url 是合作方提供给我方下载简历的地址，不是我方返回的地址。",
        "合作方只需要传 resume.file_name 与 resume.file_url，不需要传结构化简历 JSON。",
        "我方服务会使用 resume.file_url 下载简历文件，并解析为内部 resume_json，用于生成面试问题。",
        "resume.file_url 建议为带有效期的临时下载地址，有效期建议不少于 30 分钟。",
        "支持的文件类型建议为 PDF、DOC、DOCX，具体以部署环境的简历解析能力为准。",
    ])

    h(doc, "15. 字段映射层")
    p(doc, "本节是我方内部实现约定，不要求合作方调整字段名。合作方仍按 Word 文档字段传输，我方适配层负责转换。")
    table(doc, ["合作方字段", "我方内部字段/处理方式", "是否要求合作方改字段", "说明"], [
        ("request_id", "适配层保存 request_id", "否", "用于创建幂等；当前模型无该字段，需新增表/字段或 metadata 保存"),
        ("scene", "校验固定值 ai_interview", "否", "不落入核心面试模型也可以"),
        ("candidate.id", "保存为 external_candidate_id 或 metadata", "否", "当前 Candidate.candidate_id 不建议直接用外部数字 ID 覆盖"),
        ("candidate.name", "Candidate.name", "否", "直接映射"),
        ("candidate.phone", "Candidate.phone", "否", "直接映射；为空时允许空值"),
        ("candidate.email", "Candidate.email", "否", "直接映射，并用于候选人身份校验"),
        ("job.id", "保存为 external_job_id 或 metadata", "否", "用于对账和回溯合作方岗位"),
        ("job.title", "JobConfiguration.job_name / target_position", "否", "合作方叫 job.title，我方内部叫 job_name"),
        ("job.description", "JobConfiguration.job_description", "否", "合作方叫 job.description，我方内部叫 job_description"),
        ("resume.file_name", "resume_json.file_name", "否", "保存文件名"),
        ("resume.file_url", "下载简历并解析为 resume_json", "否", "这是合作方提供给我方下载简历的链接"),
        ("callback_url", "适配层保存 callback_url", "否", "面试完成后由我方 POST 回调该地址"),
        ("external_session_id", "对外返回，内部可映射 JobConfiguration.config_id", "否", "建议格式 ai_{config_id} 或直接使用 config_id"),
        ("session_url", "由我方根据 config_id 生成 interview_url", "否", "对外字段名保持 session_url"),
        ("expires_at", "JobConfiguration.expire_at", "否", "对外字段名保持 expires_at"),
    ])

    h(doc, "16. 不一致项与需要我方改造的内容")
    table(doc, ["类别", "合作方 Word 协议", "我方当前系统", "处理方式"], [
        ("创建路径", "/api/external/ai-session/create", "当前为 /api/ai-interview/hr/create-job/ 或 /api/hr/create-job/", "我方新增外部对接路径，内部复用现有创建逻辑"),
        ("创建请求结构", "嵌套结构 candidate/job/resume", "当前创建接口是扁平字段 candidate_name/job_name/resume_json", "我方适配层展开并映射，不要求合作方改字段"),
        ("岗位名称", "job.title", "job_name / target_position", "映射：job.title -> job_name、target_position"),
        ("岗位描述", "job.description", "job_description", "映射：job.description -> job_description"),
        ("简历", "resume.file_name + resume.file_url", "当前常用 resume_json 结构化简历", "我方用 file_url 下载并解析，生成内部 resume_json"),
        ("回调地址", "callback_url", "当前无外部结果回调字段", "我方新增保存 callback_url，并在面试完成后调用"),
        ("响应任务 ID", "external_session_id", "config_id / session_id", "对外返回 external_session_id，内部映射 config_id"),
        ("响应链接", "session_url", "interview_url", "对外返回 session_url，内部生成现有面试入口链接"),
        ("过期时间", "expires_at", "expire_at", "对外返回 expires_at，内部映射 expire_at"),
        ("结果结构", "result.score/summary/dimensions/report_url", "当前为 confidence_score/profile_data/report_analysis/qa_records", "我方在回调前汇总转换为合作方 result 结构"),
        ("幂等", "request_id", "当前无 request_id", "我方需要新增幂等记录"),
    ])

    h(doc, "17. 类型差异与建议")
    table(doc, ["字段", "合作方 Word 类型", "我方建议", "原因"], [
        ("candidate.id", "number/string", "保持 number/string 均可", "我方按字符串保存 external_candidate_id，避免数字和字符串差异影响对账"),
        ("job.id", "number/string", "保持 number/string 均可", "我方按字符串保存 external_job_id"),
        ("candidate.phone", "string", "建议 string", "手机号可能有前导 0 或国际区号，不建议 number"),
        ("result.score", "number", "保持 number", "我方由 confidence_score * 100 或报告分析综合得出"),
        ("result.dimensions[].score", "number", "保持 number", "我方由各轮分析分数映射"),
        ("completed_at/expires_at", "string", "保持 ISO 8601 string", "便于跨系统传输"),
    ])

    h(doc, "18. 我方有但不要求合作方传的字段")
    table(doc, ["我方字段", "处理方式", "说明"], [
        ("job_level", "使用默认值或我方内部配置", "合作方 Word 协议没有该字段，不要求传"),
        ("hard_fields", "默认空数组或我方后台配置", "合作方 Word 协议没有该字段，不要求传"),
        ("custom_questions", "默认空数组或我方后台配置", "合作方 Word 协议没有该字段，不要求传"),
        ("max_interviews", "默认 1", "合作方 Word 协议没有该字段，不要求传"),
        ("created_by", "默认 external 或 partner", "合作方 Word 协议没有该字段，不要求传"),
        ("resume_json", "由我方通过 resume.file_url 下载解析生成", "合作方不需要传结构化简历"),
    ])

    h(doc, "19. 对接流程")
    bullets(doc, [
        "合作方调用 /api/external/ai-session/create 创建 AI 面试任务，同时提供 resume.file_url 和 callback_url。",
        "我方校验参数，通过 resume.file_url 下载并解析简历，创建内部面试任务。",
        "我方返回 external_session_id 与 session_url。",
        "合作方将 session_url 发送给候选人。",
        "候选人打开 session_url，在我方系统完成身份校验并进行 AI 面试。",
        "我方生成问答记录、评分、总结和报告。",
        "我方主动调用合作方提供的 callback_url，将面试结果 POST 回传给合作方。",
        "合作方返回 received=true，我方结束该任务回调流程。",
    ])

    h(doc, "20. 安全建议")
    bullets(doc, [
        "创建任务接口使用 Authorization: Bearer token 鉴权。",
        "callback_url 建议支持签名校验，避免伪造回调。",
        "resume.file_url 建议使用短期有效的临时下载链接。",
        "双方日志中应对手机号、邮箱、简历地址等敏感信息脱敏。",
        "建议所有生产接口使用 HTTPS。",
    ])

    h(doc, "21. 开放端口前我方实现检查项")
    p(doc, "以下为我方服务开放给合作方调用前必须完成或确认的实现项，不要求合作方改变 Word 协议字段。")
    table(doc, ["检查项", "是否必须", "说明"], [
        ("新增创建任务入口 /api/external/ai-session/create", "是", "当前系统已有内部创建接口，但外部路径需要按 Word 协议新增"),
        ("Bearer Token 鉴权", "是", "请求头 Authorization: Bearer xxx 必须实际校验"),
        ("request_id 幂等存储", "是", "相同 request_id 重复请求时返回同一任务，不能重复创建"),
        ("candidate/job/resume 字段映射", "是", "保持合作方字段不变，我方内部映射"),
        ("resume.file_url 下载能力", "是", "这是合作方提供给我方下载简历的链接，需支持下载失败错误码"),
        ("简历解析为 resume_json", "是", "我方内部面试流程需要结构化简历"),
        ("返回 external_session_id/session_url/expires_at", "是", "响应字段名必须按 Word 协议返回"),
        ("保存 callback_url", "是", "面试完成后我方按该地址主动回调合作方"),
        ("结果汇总为 score/summary/dimensions/report_url", "是", "当前内部结果需转换为 Word 协议 result 结构"),
        ("回调发送与重试", "是", "0.5 秒确认窗口；未收到 received=true 时每 5 秒重试"),
        ("回调幂等", "是", "同一 request_id/external_session_id 结果不能重复写入或重复造成业务副作用"),
        ("公网端口/域名/HTTPS", "上线建议", "开放端口前确认防火墙、安全组、反向代理、HTTPS 证书"),
    ])

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
