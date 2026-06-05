from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "AI面试任务创建与结果查询API对接文档-当前系统版.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)


def add_code(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F7F7")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    for index, line in enumerate(code.strip().splitlines()):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(8.5)


def add_field_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["字段", "类型", "必填", "说明"]
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
        set_cell_shading(table.rows[0].cells[i], "EAF2F8")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    doc.add_paragraph()


def add_kv_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for i, header in enumerate(["项目", "内容"]):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
        set_cell_shading(table.rows[0].cells[i], "EAF2F8")
    for key, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], key)
        set_cell_text(cells[1], value)
    doc.add_paragraph()


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        if level == 1:
            run.font.color.rgb = RGBColor(31, 78, 121)
        else:
            run.font.color.rgb = RGBColor(47, 84, 150)
    return paragraph


def add_paragraph(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(10.5)


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AI 面试任务创建与结果查询 API 对接文档")
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("当前系统版")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(90, 90, 90)

    add_heading(doc, "1. 接口概述")
    add_paragraph(
        doc,
        "本接口用于在当前 AI 面试系统中创建候选人的面试任务，生成可发送给候选人的面试链接。"
        "候选人通过链接完成身份校验并进入 AI 面试。面试完成后，系统将问答记录、画像数据、综合可信度、"
        "报告分析结果等数据存储在本系统数据库中，HR 可通过报告接口查询或触发分析。"
    )
    add_paragraph(
        doc,
        "当前系统不是第三方回调模式：创建任务后返回本系统面试链接，结果通过本系统报告接口查询；"
        "如果后续需要第三方回调，可在本协议基础上新增回调适配层。"
    )

    add_heading(doc, "2. 基础路径与鉴权")
    add_kv_table(
        doc,
        [
            ("服务地址", "由部署环境决定，例如：http://localhost:8000"),
            ("主要 API 前缀", "/api/ai-interview/"),
            ("兼容创建路径", "/api/hr/create-job/"),
            ("请求格式", "application/json；简历解析接口使用 multipart/form-data"),
            ("鉴权方式", "当前代码未内置 Bearer Token 鉴权，生产环境建议在网关或视图层增加鉴权"),
        ],
    )

    add_heading(doc, "3. 创建 AI 面试任务")
    add_kv_table(
        doc,
        [
            ("请求方式", "POST"),
            ("推荐路径", "/api/ai-interview/hr/create-job/"),
            ("兼容路径", "/api/hr/create-job/"),
            ("Content-Type", "application/json"),
        ],
    )
    add_code(
        doc,
        r'''
{
  "job_name": "AI应用工程师",
  "job_description": "岗位详情，包含职位描述、岗位职责、任职要求、加分项等完整信息。",
  "job_level": "初级",
  "hard_fields": ["英语水平", "是否接受出差"],
  "custom_questions": ["会开车吗"],
  "max_interviews": 1,
  "days_valid": 3,
  "candidate_name": "张新",
  "candidate_phone": "17337075112",
  "candidate_email": "2601413168@qq.com",
  "created_by": "admin",
  "resume_json": {
    "name": "张新",
    "phone": "17337075112",
    "email": "2601413168@qq.com",
    "file_name": "张新的简历.pdf",
    "summary": "候选人简历摘要",
    "education": [],
    "work_experience": [],
    "projects": [],
    "skills": []
  }
}
''',
    )
    add_field_table(
        doc,
        [
            ("job_name", "string", "是", "岗位名称，同时作为 target_position 存储"),
            ("job_description", "string", "否", "岗位描述/JD，用于生成面试问题"),
            ("job_level", "string", "否", "岗位级别，默认使用系统默认值"),
            ("hard_fields", "array", "否", "必问硬性字段列表，例如学历、证书、驾驶能力等"),
            ("custom_questions", "array/object", "否", "自定义面试问题；当前前端常用数组格式"),
            ("max_interviews", "number", "否", "请求字段保留；当前创建逻辑实际固定存储为 1"),
            ("days_valid", "number", "否", "面试链接有效天数，默认 3 天，最小 1 天"),
            ("candidate_name", "string", "是", "候选人姓名"),
            ("candidate_phone", "string", "否", "候选人手机号"),
            ("candidate_email", "string", "是", "候选人邮箱"),
            ("created_by", "string", "否", "创建人标识，默认 system"),
            ("resume_json", "object", "否", "已解析的结构化简历 JSON；为空时面试开始阶段会使用候选人基础信息兜底"),
        ],
    )

    add_heading(doc, "4. 创建任务响应")
    add_code(
        doc,
        r'''
{
  "config_id": "4089dfa7fbdb42ce",
  "job_name": "AI应用工程师",
  "interview_url": "http://localhost:8000/api/interview/welcome/?config_id=4089dfa7fbdb42ce",
  "expire_at": "2026-06-05T10:00:00+08:00",
  "max_interviews": 1,
  "message": "岗位配置创建成功"
}
''',
    )
    add_field_table(
        doc,
        [
            ("config_id", "string", "是", "本系统生成的面试任务配置 ID，也是候选人入口链接的核心参数"),
            ("job_name", "string", "是", "岗位名称"),
            ("interview_url", "string", "是", "候选人访问链接，HR 可发送给候选人"),
            ("expire_at", "string/null", "否", "链接过期时间，ISO 8601 格式"),
            ("max_interviews", "number", "是", "最大面试次数；当前响应返回请求值"),
            ("message", "string", "是", "创建结果说明"),
        ],
    )

    add_heading(doc, "5. 简历解析接口")
    add_paragraph(doc, "如调用方只有简历文件，建议先使用简历解析接口生成 resume_json，再调用创建任务接口。")
    add_kv_table(
        doc,
        [
            ("请求方式", "POST"),
            ("路径", "/api/ai-interview/resume/parse/ 或 /api/resume/parse/"),
            ("Content-Type", "multipart/form-data"),
            ("文件字段", "resume"),
        ],
    )
    add_paragraph(doc, "解析结果会包含 basic_info、education、work_experience、projects、skills、summary、file_name、parsed_at 等结构化字段。")

    add_heading(doc, "6. 查询岗位任务信息")
    add_kv_table(
        doc,
        [
            ("请求方式", "GET"),
            ("路径", "/api/interview/job-info/?config_id={config_id}"),
            ("用途", "候选人进入欢迎页时查询岗位名称、级别和剩余面试次数"),
        ],
    )
    add_code(
        doc,
        r'''
{
  "config_id": "4089dfa7fbdb42ce",
  "job_name": "AI应用工程师",
  "job_level": "初级",
  "target_position": "AI应用工程师",
  "remaining_interviews": 1
}
''',
    )

    add_heading(doc, "7. 候选人身份校验并开始面试")
    add_kv_table(
        doc,
        [
            ("请求方式", "POST"),
            ("路径", "/api/interview/verify-and-start/"),
            ("Content-Type", "application/json"),
        ],
    )
    add_code(
        doc,
        r'''
{
  "config_id": "4089dfa7fbdb42ce",
  "name": "张新",
  "email": "2601413168@qq.com"
}
''',
    )
    add_code(
        doc,
        r'''
{
  "session_id": "6c274738-8727-42c1-894b-c9d353ae6d12",
  "candidate_id": "张新_2601413168@qq.com",
  "question": "",
  "round_number": 0,
  "interview_url": "http://localhost:8000/start-interview/?session_id=..."
}
''',
    )

    add_heading(doc, "8. 面试会话接口")
    add_kv_table(
        doc,
        [
            ("页面入口", "GET /start-interview/?session_id={session_id}&candidate_id={candidate_id}&config_id={config_id}"),
            ("会话 API", "POST /api/ai-interview/session/"),
            ("会话详情", "GET /api/ai-interview/session/{session_id}/"),
            ("释放锁", "POST /api/interview/release-lock/"),
        ],
    )
    add_paragraph(doc, "当前系统支持文本和语音面试流程。语音能力主要通过 WebSocket 与语音服务实现，接口调用方通常只需要使用创建任务与候选人入口链接。")

    add_heading(doc, "9. 面试结果查询")
    add_kv_table(
        doc,
        [
            ("请求方式", "GET"),
            ("路径", "/api/ai-interview/hr/candidates/interview-report/data/?candidate_id={candidate_id}&config_id={config_id}"),
            ("用途", "查询候选人问答记录、完成状态、综合可信度和已有分析结果"),
        ],
    )
    add_code(
        doc,
        r'''
{
  "candidate_id": "张新_2601413168@qq.com",
  "candidate_name": "张新",
  "interview_completed": true,
  "confidence_score": 0.86,
  "report_analyzed": true,
  "updated_at": "2026-06-02T15:30:00+08:00",
  "qa_records": [
    {
      "round": 1,
      "round_type": "hard_field",
      "question": "请介绍一下你的项目经历",
      "answer": "我主要负责...",
      "score": 86,
      "analysis_status": "analyzed"
    }
  ]
}
''',
    )
    add_field_table(
        doc,
        [
            ("candidate_id", "string", "是", "候选人唯一标识，当前格式通常为 姓名_邮箱"),
            ("candidate_name", "string", "是", "候选人姓名"),
            ("interview_completed", "boolean", "是", "面试是否完成"),
            ("confidence_score", "number", "是", "综合可信度，0-1 区间；展示时通常乘以 100"),
            ("report_analyzed", "boolean", "是", "是否已经生成问答分析"),
            ("updated_at", "string/null", "否", "最近更新时间"),
            ("qa_records", "array", "是", "问答记录及分析结果"),
        ],
    )

    add_heading(doc, "10. 触发报告分析")
    add_kv_table(
        doc,
        [
            ("请求方式", "POST"),
            ("路径", "/api/ai-interview/hr/candidates/interview-report/analyze/"),
            ("Content-Type", "application/json"),
            ("用途", "对面试问答记录进行质量分析，并写入 profile_data.report_analysis"),
        ],
    )
    add_code(
        doc,
        r'''
{
  "candidate_id": "张新_2601413168@qq.com",
  "config_id": "4089dfa7fbdb42ce"
}
''',
    )
    add_code(
        doc,
        r'''
{
  "success": true,
  "message": "问答报告分析完成",
  "candidate_id": "张新_2601413168@qq.com",
  "candidate_name": "张新",
  "interview_completed": true,
  "confidence_score": 0.86,
  "report_analyzed": true,
  "qa_records": [],
  "qa_count": 6
}
''',
    )

    add_heading(doc, "11. 数据库存储结构")
    add_field_table(
        doc,
        [
            ("Candidate.candidate_id", "string", "是", "候选人唯一标识，默认由姓名和邮箱拼接"),
            ("Candidate.name/phone/email", "string", "是/否", "候选人基础信息"),
            ("Candidate.interview_id", "string", "否", "关联的 config_id"),
            ("JobConfiguration.config_id", "string", "是", "面试任务配置 ID"),
            ("JobConfiguration.job_name/job_description/job_level", "string", "是/否", "岗位基础信息"),
            ("JobConfiguration.hard_fields", "JSON array", "否", "硬性问题字段"),
            ("JobConfiguration.custom_questions", "JSON", "否", "自定义问题"),
            ("JobConfiguration.resume_json", "JSON object", "否", "结构化简历数据"),
            ("TalentProfile.session_id", "string", "是", "面试会话 ID"),
            ("TalentProfile.raw_resume", "JSON object", "是", "面试时使用的简历数据"),
            ("TalentProfile.job_config", "JSON object", "是", "面试时使用的岗位配置快照"),
            ("TalentProfile.dialogue_history", "JSON array", "是", "原始对话记录"),
            ("TalentProfile.profile_data", "JSON object", "是", "画像、问答记录、报告分析等完整结果"),
            ("TalentProfile.confidence_score", "number", "是", "综合可信度"),
            ("TalentProfile.interview_completed", "boolean", "是", "面试完成状态"),
        ],
    )

    add_heading(doc, "12. 状态与错误")
    add_paragraph(doc, "当前系统未集中定义统一 error_code，错误响应通常使用 error 字段返回中文错误信息。")
    add_code(
        doc,
        r'''
{
  "error": "config_id不能为空"
}
''',
    )
    add_field_table(
        doc,
        [
            ("400", "HTTP status", "否", "请求参数缺失、身份信息不匹配、链接过期或面试已完成"),
            ("404", "HTTP status", "否", "配置、候选人或面试记录不存在"),
            ("423", "HTTP status", "否", "面试链接正在使用中，锁定状态"),
            ("500", "HTTP status", "否", "系统内部错误"),
        ],
    )

    add_heading(doc, "13. 当前系统与外部回调协议的差异")
    add_bullets(
        doc,
        [
            "当前系统创建任务时接收 resume_json，即结构化简历；外部回调协议通常只传 resume.file_name 与 resume.file_url。",
            "当前系统返回 interview_url 与 config_id；外部协议常见字段为 session_url 与 external_session_id。",
            "当前系统没有 /api/integrations/ai-session/callback/ 回调入口，面试结果通过报告查询接口读取。",
            "当前系统结果以 dialogue_history、qa_records、report_analysis、confidence_score 存储；外部协议常见 result.score、summary_for_hr、dimensions、report_url 需要额外汇总映射。",
            "当前系统未内置 request_id 幂等字段；如需要对接第三方任务创建协议，建议新增 request_id、callback_url、external_status 等字段。",
        ],
    )

    add_heading(doc, "14. 推荐对接流程")
    add_bullets(
        doc,
        [
            "调用方上传简历文件到本系统，使用简历解析接口生成 resume_json。",
            "调用方调用创建任务接口，传入岗位、候选人和 resume_json。",
            "系统返回 config_id 与 interview_url。",
            "HR 将 interview_url 发送给候选人。",
            "候选人打开链接，完成姓名和邮箱校验后进入 AI 面试。",
            "面试完成后，HR 或调用方通过报告数据接口查询 interview_completed、confidence_score 和 qa_records。",
            "如 report_analyzed 为 false，可调用报告分析接口生成逐轮分析结果。",
        ],
    )

    add_heading(doc, "15. 安全建议")
    add_bullets(
        doc,
        [
            "生产环境建议为创建任务、报告查询、报告分析接口增加 Bearer Token、签名或网关鉴权。",
            "不要在 URL 中暴露过多候选人敏感信息；当前候选人入口主要使用 config_id。",
            "简历文件和结构化简历数据属于敏感信息，应限制访问权限并做好日志脱敏。",
            "面试链接建议保留过期时间，days_valid 不宜过长。",
            "如果未来接入第三方回调，应增加 request_id 幂等校验和回调签名校验。",
        ],
    )

    section = doc.add_section(WD_SECTION.CONTINUOUS)
    section.footer.paragraphs[0].text = "AI 面试系统 API 对接文档 - 当前系统版"
    for run in section.footer.paragraphs[0].runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(8)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_doc())
